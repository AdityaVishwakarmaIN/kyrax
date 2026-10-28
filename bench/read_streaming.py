"""Opt-in first-row latency and bounded-consumption comparisons.

All libraries perform the same Python checksum/type bookkeeping. Calamine
materializes its complete sheet; the two other readers consume row iterators.
Source revision is explicit because the package may be a frozen source copy.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib.metadata
import itertools
import json
import os
import platform
import resource
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from fixtures import SIZES, ensure_fixture, iter_rows

HERE = Path(__file__).resolve().parent
LIBS = ("kyrax", "openpyxl-ro", "calamine")
TASKS = ("read_first", "read_stream")


def _cpu():
    usage = resource.getrusage(resource.RUSAGE_SELF)
    return usage.ru_utime + usage.ru_stime


def _rss():
    for line in Path("/proc/self/status").read_text().splitlines():
        if line.startswith("VmHWM:"):
            return int(line.split()[1]) / 1024
    raise RuntimeError("Linux VmHWM is unavailable")


def _canonical(value):
    # Include type families: bool/int equality must not conceal lost types.
    # This fixture contains numeric values well within double precision.
    if isinstance(value, bool):
        return "bool", value
    if isinstance(value, dt.datetime):
        return "date", value.toordinal(), value.hour, value.minute, value.second, value.microsecond
    if isinstance(value, dt.date):
        return "date", value.toordinal(), 0, 0, 0, 0
    if isinstance(value, (int, float)):
        return "number", float(value)
    if isinstance(value, str):
        return "string", value
    if value is None:
        return "blank", None
    return type(value).__name__, repr(value)


def consume(rows, limit=None):
    """Consume without retaining rows; checksum covers every consumed cell."""
    digest = hashlib.blake2b(digest_size=16)
    observed_types = []
    count = 0
    if limit is not None:
        rows = itertools.islice(rows, limit)
    for row in rows:
        if len(observed_types) < len(row):
            observed_types.extend(set() for _ in range(len(row) - len(observed_types)))
        for column, value in enumerate(row):
            observed_types[column].add(type(value).__name__)
        digest.update(repr(tuple(_canonical(value) for value in row)).encode("utf-8"))
        digest.update(b"\n")
        count += 1
    return {
        "count": count,
        "checksum": digest.hexdigest(),
        "read_types": [sorted(types) for types in observed_types],
    }


def _worker(task, lib, size, fixture, output, package_root):
    sys.path.insert(0, package_root)
    limit = 10 if task == "read_first" else None
    c0, w0 = _cpu(), time.perf_counter()
    if lib == "kyrax":
        import kyrax

        imported_path = str(Path(kyrax.__file__).resolve())
        if not Path(imported_path).is_relative_to(Path(package_root).resolve()):
            raise RuntimeError(f"Imported Kyrax outside requested package root: {imported_path}")
        workbook = kyrax.load_workbook(fixture, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    elif lib == "openpyxl-ro":
        import openpyxl

        imported_path = str(Path(openpyxl.__file__).resolve())
        workbook = openpyxl.load_workbook(fixture, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    else:
        from python_calamine import CalamineWorkbook

        imported_path = "python_calamine"
        workbook = CalamineWorkbook.from_path(fixture)
        rows = workbook.get_sheet_by_index(0).to_python()
    actual = consume(rows, limit)
    if hasattr(rows, "close"):
        rows.close()
    if hasattr(workbook, "close"):
        workbook.close()
    cpu, wall = _cpu() - c0, time.perf_counter() - w0
    peak = _rss()  # Capture before untimed source generation/verification.
    expected = consume(iter_rows(SIZES[size]), limit)
    ok = actual["count"] == expected["count"] and actual["checksum"] == expected["checksum"]
    result = {
        "python_version": sys.version.split()[0],
        "cpu_task_s": cpu,
        "wall_task_s": wall,
        "rss_mb": peak,
        "verify": "ok (count and typed full-cell checksum)" if ok else "FAIL typed checksum/count",
        "expected_checksum": expected["checksum"],
        "imported_path": imported_path,
        "library_version": importlib.metadata.version(
            {"kyrax": "kyrax", "openpyxl-ro": "openpyxl", "calamine": "python-calamine"}[lib]
        )
        if lib != "kyrax"
        else "source package (see extension SHA256)",
        **actual,
    }
    Path(output).write_text(json.dumps(result))


def _summary(values):
    return {"median": statistics.median(values), "min": min(values), "max": max(values)}


def _render(data):
    meta = data["meta"]
    lines = [
        "# Partial and streaming read benchmark",
        "",
        f"Source: `{meta['source_revision']}` · {meta['date_utc']} · Python {meta['python']}",
        f"Package root: `{meta['package_root']}`",
        "",
        "CPU/wall include import, open, identical Python checksum/type bookkeeping, and close.",
        "Peak RSS is captured before untimed expected-data validation. First read consumes ten",
        "rows including the header. Calamine eagerly loads all rows in both tasks. Shared",
        "strings may load eagerly in every reader. FAIL rows have unequal type semantics.",
        "",
        "| task | size | library | CPU s | wall s | peak MiB | n | verification |",
        "|---|---|---|---:|---:|---:|---:|---|",
    ]
    for row in data["rows"]:
        lines.append(
            f"| {row['task']} | {row['size']} | {row['lib']} | "
            f"{row['cpu_task_s']['median']:.4f} | {row['wall_task_s']['median']:.4f} | "
            f"{row['rss_mb']['median']:.1f} | {row['n']} | {row['verify']} |"
        )
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", nargs=6, metavar="ARG", help=argparse.SUPPRESS)
    parser.add_argument("--package-root", type=Path)
    parser.add_argument("--source-revision")
    parser.add_argument("--python", default=sys.executable)
    parser.add_argument("--cache", type=Path, default=HERE / ".cache")
    parser.add_argument("--out", type=Path, default=HERE / "results" / "sprint_read_streaming")
    parser.add_argument("--sizes", nargs="+", choices=SIZES, default=list(SIZES))
    parser.add_argument("--tasks", nargs="+", choices=TASKS, default=list(TASKS))
    parser.add_argument("--libs", nargs="+", choices=LIBS, default=list(LIBS))
    parser.add_argument("--repeats", type=int, default=3)
    args = parser.parse_args()
    if args.worker:
        _worker(*args.worker)
        return
    if args.package_root is None or args.source_revision is None:
        parser.error("--package-root and --source-revision are required for explicit provenance")
    package_root = args.package_root.resolve()
    if not (package_root / "kyrax" / "__init__.py").is_file():
        parser.error("--package-root must contain kyrax/__init__.py")
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    fixtures = {size: ensure_fixture(str(args.cache), size) for size in args.sizes}
    rows = []
    env = dict(os.environ, PYTHONHASHSEED="0")
    with tempfile.TemporaryDirectory(prefix="kyrax-stream-bench-") as directory:
        result_path = str(Path(directory) / "result.json")
        for task in args.tasks:
            for size in args.sizes:
                measurements = {lib: [] for lib in args.libs}
                for repeat in range(args.repeats):
                    for lib in args.libs:
                        command = [
                            args.python,
                            str(Path(__file__).resolve()),
                            "--worker",
                            task,
                            lib,
                            size,
                            fixtures[size],
                            result_path,
                            str(package_root),
                        ]
                        completed = subprocess.run(command, env=env, capture_output=True, text=True)
                        if completed.returncode:
                            raise RuntimeError(f"{task}/{size}/{lib}: {completed.stderr[-4000:]}")
                        result = json.loads(Path(result_path).read_text())
                        measurements[lib].append(result)
                        print(
                            f"{task}/{size}/{lib} rep{repeat}: "
                            f"CPU={result['cpu_task_s']:.3f}s RSS={result['rss_mb']:.1f}MiB "
                            f"{result['verify']}",
                            file=sys.stderr,
                            flush=True,
                        )
                for lib, samples in measurements.items():
                    rows.append(
                        {
                            "task": task,
                            "size": size,
                            "lib": lib,
                            "n": len(samples),
                            **{
                                key: _summary([sample[key] for sample in samples])
                                for key in ("cpu_task_s", "wall_task_s", "rss_mb")
                            },
                            "verify": (
                                "ok (count and typed full-cell checksum)"
                                if all(s["verify"].startswith("ok") for s in samples)
                                else "FAIL typed checksum/count"
                            ),
                            "samples": samples,
                        }
                    )
    extensions = sorted((package_root / "kyrax").glob("*.so"))
    data = {
        "meta": {
            "date_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
            "source_revision": args.source_revision,
            "package_root": str(package_root),
            "extension_sha256": {
                p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in extensions
            },
            "python": rows[0]["samples"][0]["python_version"],
            "worker_interpreter": args.python,
            "platform": platform.platform(),
            "semantics": (
                "first10/all rows; identical timed typed checksum consumer; calamine eager"
            ),
        },
        "rows": rows,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.with_suffix(".json").write_text(json.dumps(data, indent=2))
    args.out.with_suffix(".md").write_text(_render(data))
    print(_render(data))


if __name__ == "__main__":
    main()
