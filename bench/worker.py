"""One measurement in a fresh process: ``worker.py TASK LIB SIZE FIXTURE OUT RESULT``.

Timed section = library import + the task itself (interpreter start-up and
input-data construction are excluded; they are identical for every library).
The parent additionally records whole-process CPU and peak RSS via wait4().
"""

from __future__ import annotations

import datetime as dt
import json
import os
import resource
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fixtures import NCOLS, SIZES, iter_rows, make_rows  # noqa: E402

_READ_RESULT = None

EDIT_CELLS = 1000  # cells overwritten in column B + cells added in a new column I


def _vm_hwm_mb() -> float:
    """Peak RSS of this process image (VmHWM). Unlike the parent's wait4
    ru_maxrss it is not polluted by the parent's RSS at fork/exec time."""
    with open("/proc/self/status") as f:
        for line in f:
            if line.startswith("VmHWM:"):
                return int(line.split()[1]) / 1024.0
    return float("nan")


def _cpu() -> float:
    r = resource.getrusage(resource.RUSAGE_SELF)
    return r.ru_utime + r.ru_stime


# --------------------------------------------------------------------- IMPORT
def _import(mod):
    def fn(fx, out, rows):
        __import__(mod)
        return 0, 0

    return fn


# ---------------------------------------------------------------- READ values
def read_values_kyrax(fx, out, rows):
    import kyrax

    wb = kyrax.load_workbook(fx, read_only=True)
    data = list(wb.active.iter_rows(values_only=True))
    global _READ_RESULT
    _READ_RESULT = data
    return len(data), len(data[0])


def read_values_calamine(fx, out, rows):
    from python_calamine import CalamineWorkbook

    data = CalamineWorkbook.from_path(fx).get_sheet_by_index(0).to_python()
    global _READ_RESULT
    _READ_RESULT = data
    return len(data), len(data[0])


def read_values_openpyxl(fx, out, rows):
    import openpyxl

    wb = openpyxl.load_workbook(fx, read_only=True)
    data = list(wb.active.iter_rows(values_only=True))
    wb.close()
    global _READ_RESULT
    _READ_RESULT = data
    return len(data), len(data[0])


# ------------------------------------------------------------- READ DataFrame
def read_df_kyrax(fx, out, rows):
    import kyrax

    df = kyrax.read_excel(fx).load_sheet(0).to_pandas()
    global _READ_RESULT
    _READ_RESULT = df
    return len(df) + 1, len(df.columns)


def read_df_pandas_calamine(fx, out, rows):
    import pandas as pd

    df = pd.read_excel(fx, engine="calamine")
    global _READ_RESULT
    _READ_RESULT = df
    return len(df) + 1, len(df.columns)


def read_df_pandas_openpyxl(fx, out, rows):
    import pandas as pd

    df = pd.read_excel(fx, engine="openpyxl")
    global _READ_RESULT
    _READ_RESULT = df
    return len(df) + 1, len(df.columns)


# ---------------------------------------------------------------------- WRITE
def write_noop(fx, out, rows):
    """Floor: data construction only (shows the input-data RSS baseline)."""
    return len(rows), len(rows[0])


def write_kyrax(fx, out, rows):
    import kyrax

    kyrax.write_excel_turbo(out, [{"name": "Sheet1", "rows": rows}])
    return len(rows), len(rows[0])


def write_kyrax_stream(fx, out, rows):
    import kyrax

    kyrax.write_excel_turbo_stream(out, [{"name": "Sheet1", "rows": rows}])
    return len(rows), len(rows[0])


def _xlsxwriter(out, rows, constant_memory):
    import xlsxwriter

    wb = xlsxwriter.Workbook(
        out, {"default_date_format": "yyyy-mm-dd", "constant_memory": constant_memory}
    )
    ws = wb.add_worksheet("Sheet1")
    write_row = ws.write_row
    for r, row in enumerate(rows):
        write_row(r, 0, row)
    wb.close()
    return len(rows), len(rows[0])


def write_xlsxwriter(fx, out, rows):
    return _xlsxwriter(out, rows, False)


def write_xlsxwriter_cm(fx, out, rows):
    return _xlsxwriter(out, rows, True)


def _openpyxl_write(out, rows, write_only):
    import openpyxl

    wb = openpyxl.Workbook(write_only=write_only)
    ws = wb.create_sheet("Sheet1") if write_only else wb.active
    append = ws.append
    for row in rows:
        append(row)
    wb.save(out)
    return len(rows), len(rows[0])


def write_openpyxl(fx, out, rows):
    return _openpyxl_write(out, rows, False)


def write_openpyxl_wo(fx, out, rows):
    return _openpyxl_write(out, rows, True)


# ------------------------------------------------------------ GENERATOR WRITE
def write_generator_noop(fx, out, rows):
    count = sum(1 for _ in rows)
    return count, NCOLS


def write_generator_kyrax(fx, out, rows):
    import kyrax

    kyrax.write_excel_turbo_stream(out, [{"name": "Sheet1", "rows": rows}])
    return int(os.environ["BENCH_NROWS"]) + 1, NCOLS


def write_generator_xlsxwriter(fx, out, rows):
    import xlsxwriter

    wb = xlsxwriter.Workbook(out, {"default_date_format": "yyyy-mm-dd", "constant_memory": True})
    ws = wb.add_worksheet("Sheet1")
    count = 0
    for count, row in enumerate(rows, 1):
        ws.write_row(count - 1, 0, row)
    wb.close()
    return count, NCOLS


def write_generator_openpyxl(fx, out, rows):
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    count = 0
    for count, row in enumerate(rows, 1):
        ws.append(row)
    wb.save(out)
    return count, NCOLS


def verify_read(task, nrows):
    """Validate outside timing; numeric representation may be int or float.

    Boolean, string and date types are checked separately: Python's bool/int
    equality and Excel date serials must not conceal reader type loss.
    """
    if task == "read_values":
        data = _READ_RESULT
        expected = iter_rows(nrows)
        if len(data) != nrows + 1:
            return f"FAIL rows {len(data)} != {nrows + 1}", {}
    elif task == "read_df":
        data = _READ_RESULT
        if len(data) != nrows or len(data.columns) != NCOLS:
            return f"FAIL DataFrame shape {data.shape}", {}
        if list(data.columns) != next(iter_rows(nrows)):
            return "FAIL DataFrame columns differ", {}
        data = data.itertuples(index=False, name=None)
        expected = iter_rows(nrows)
        next(expected)
    else:
        return "n/a", {}
    types = [set() for _ in range(NCOLS)]
    for ri, (actual, exp) in enumerate(zip(data, expected)):
        if len(actual) != len(exp):
            return f"FAIL row {ri} width {len(actual)} != {len(exp)}", {}
        for ci, (got, want) in enumerate(zip(actual, exp)):
            types[ci].add(type(got).__name__)
            if isinstance(want, bool):
                ok = isinstance(got, bool) and got == want
            elif isinstance(want, dt.datetime):
                ok = isinstance(got, (dt.date, dt.datetime))
                if ok and isinstance(got, dt.datetime):
                    ok = got.time() == want.time() and got.date() == want.date()
                elif ok:
                    ok = got == want.date()
            elif isinstance(want, (int, float)):
                ok = isinstance(got, (int, float)) and not isinstance(got, bool) and got == want
            else:
                ok = type(got) is type(want) and got == want
            if not ok:
                return (f"FAIL row {ri} col {ci}: {got!r} ({type(got).__name__}) != "
                        f"{want!r} ({type(want).__name__})"), {}
    return "ok (all values; bool/date/numeric types)", {
        str(i): sorted(t) for i, t in enumerate(types)
    }


# ----------------------------------------------------------------------- EDIT
def _edits(nrows):
    k = min(EDIT_CELLS, nrows)
    return [(r, r * 0.5, f"edited-{r}") for r in range(2, k + 2)]


def edit_kyrax(fx, out, rows):
    import kyrax

    nrows = int(os.environ["BENCH_NROWS"])
    wb = kyrax.load_workbook(fx)
    ws = wb.active
    for r, num, text in _edits(nrows):
        ws.cell(row=r, column=2, value=num)
        ws.cell(row=r, column=NCOLS + 1, value=text)
    wb.save(out)
    return nrows + 1, NCOLS + 1


def edit_openpyxl(fx, out, rows):
    import openpyxl

    nrows = int(os.environ["BENCH_NROWS"])
    wb = openpyxl.load_workbook(fx)
    ws = wb.active
    for r, num, text in _edits(nrows):
        ws.cell(row=r, column=2, value=num)
        ws.cell(row=r, column=NCOLS + 1, value=text)
    wb.save(out)
    return nrows + 1, NCOLS + 1


TASKS = {
    "import": {
        "kyrax": _import("kyrax"),
        "calamine": _import("python_calamine"),
        "openpyxl": _import("openpyxl"),
        "xlsxwriter": _import("xlsxwriter"),
        "pandas": _import("pandas"),
    },
    "read_values": {
        "kyrax": read_values_kyrax,
        "calamine": read_values_calamine,
        "openpyxl-ro": read_values_openpyxl,
    },
    "read_df": {
        "kyrax": read_df_kyrax,
        "pandas-calamine": read_df_pandas_calamine,
        "pandas-openpyxl": read_df_pandas_openpyxl,
    },
    "write": {
        "(data only)": write_noop,
        "kyrax": write_kyrax,
        "kyrax-stream": write_kyrax_stream,
        "xlsxwriter": write_xlsxwriter,
        "xlsxwriter-cm": write_xlsxwriter_cm,
        "openpyxl": write_openpyxl,
        "openpyxl-wo": write_openpyxl_wo,
    },
    "write_stream": {
        "(data only)": write_generator_noop,
        "kyrax-stream": write_generator_kyrax,
        "xlsxwriter-cm": write_generator_xlsxwriter,
        "openpyxl-wo": write_generator_openpyxl,
    },
    "edit": {
        "kyrax": edit_kyrax,
        "openpyxl": edit_openpyxl,
    },
}


def main() -> None:
    task, lib, size, fx, out, result = sys.argv[1:7]
    fn = TASKS[task][lib]
    rows = make_rows(SIZES[size]) if task == "write" else None
    if task == "write_stream":
        rows = iter_rows(SIZES[size])
    os.environ["BENCH_NROWS"] = str(SIZES[size])
    c0, w0 = _cpu(), time.perf_counter()
    shape = fn(fx, out, rows)
    c1, w1 = _cpu(), time.perf_counter()
    # Snapshot peak memory before untimed validation constructs expected rows.
    peak_rss = _vm_hwm_mb()
    verdict, value_types = verify_read(task, SIZES[size])
    with open(result, "w") as f:
        json.dump(
            {"cpu_task": c1 - c0, "wall_task": w1 - w0, "rss_mb": peak_rss,
             "verify_read": verdict, "read_types": value_types,
             "shape": list(shape)},
            f,
        )


if __name__ == "__main__":
    main()
