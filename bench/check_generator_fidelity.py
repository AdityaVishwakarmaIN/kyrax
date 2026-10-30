"""Independent saved-output checks; run with PYTHONPATH=python in the benchmark venv."""

import datetime as dt
import json
import tempfile
from pathlib import Path

import kyrax
import openpyxl
from python_calamine import CalamineWorkbook


def normalized(value):
    if isinstance(value, dt.datetime) and value.time() == dt.time():
        return value.date()
    return value


def check(actual, expected):
    assert len(actual) == len(expected), (actual, expected)
    for row_actual, row_expected in zip(actual, expected):
        assert len(row_actual) == len(row_expected), (row_actual, row_expected)
        for got, want in zip(row_actual, row_expected):
            got, want = normalized(got), normalized(want)
            assert got == want, (got, want)
            if isinstance(want, bool):
                assert type(got) is bool, (got, want)
            elif isinstance(want, (str, dt.date, dt.time, dt.timedelta)):
                assert type(got) is type(want), (got, want)


results = []
with tempfile.TemporaryDirectory() as directory:
    for epoch in [False, True]:
        for mode in ["inline", "sst"]:
            path = Path(directory) / f"generator-{epoch}-{mode}.xlsx"
            header = ["id", "bool", "text", "date", "time", "duration", "datetime", "styled"]
            values = [
                [
                    1,
                    True,
                    "007",
                    dt.date(2024, 1, 2),
                    dt.time(1, 2, 3),
                    dt.timedelta(days=2, seconds=5),
                    dt.datetime(2024, 2, 29, 12, 30, 15),
                    99,
                ],
                [
                    2,
                    False,
                    "next",
                    dt.date(1904, 1, 2),
                    dt.time(12),
                    dt.timedelta(days=3),
                    dt.datetime(2020, 1, 1, 6),
                    100,
                ],
            ]

            def source():
                yield header
                for row in values:
                    yield [*row[:-1], {"value": row[-1], "style": {"font": {"bold": True}}}]

            kyrax.write_excel_turbo_stream(
                path,
                [
                    {"name": "Empty", "rows": iter(())},
                    {"name": "Data", "rows": source()},
                    {"name": "Sparse", "rows": iter([[], [None, 7], [None, None, 8]])},
                ],
                date1904=epoch,
                string_mode=mode,
                active_tab=1,
            )
            for read_only in [False, True]:
                workbook = openpyxl.load_workbook(path, read_only=read_only, data_only=True)
                assert workbook.sheetnames == ["Empty", "Data", "Sparse"]
                assert workbook.active.title == "Data"
                check(list(workbook["Data"].values), [header, *values])
                assert workbook["Data"].cell(2, 8).font.bold
                assert workbook["Sparse"].cell(2, 2).value == 7
                assert workbook["Sparse"].cell(3, 3).value == 8
                workbook.close()
            workbook = CalamineWorkbook.from_path(str(path))
            check(workbook.get_sheet_by_name("Data").to_python(), [header, *values])
            results.append(
                {
                    "date1904": epoch,
                    "string_mode": mode,
                    "openpyxl_normal": "pass",
                    "openpyxl_read_only": "pass",
                    "python_calamine": "pass",
                }
            )
print(json.dumps({"generator_fidelity": results}, indent=2))
