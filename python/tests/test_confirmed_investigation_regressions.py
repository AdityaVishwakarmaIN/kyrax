"""Focused regressions for the confirmed investigation findings."""

from __future__ import annotations

import ast
import io
from pathlib import Path

import kyrax
import kyrax.pandas  # noqa: F401 - register the pandas engine
import openpyxl
import pytest
from kyrax.formatting.rule import Rule
from openpyxl.formatting.rule import FormulaRule as OpenpyxlFormulaRule
from openpyxl.worksheet.datavalidation import DataValidation as OpenpyxlDataValidation


def _write(sheets, **kwargs):
    return kyrax.write_excel_turbo_bytes(sheets, **kwargs)


def _rows(**options):
    return _write([{"name": "S", "rows": [["A", "B"], [1, 2]], **options}])


def test_grid_rejects_metadata_and_hidden_state():
    np = pytest.importorskip("numpy")
    for key, value in [
        ("merges", ["A1:B1"]),
        ("auto_filter", "A1:B2"),
        ("tables", [{"name": "T", "ref": "A1:B2", "columns": ["A", "B"]}]),
        ("comments", [{"ref": "A1", "text": "x"}]),
        ("charts", [{"type": "line", "series": [{"val": "S!$A$1:$A$2"}]}]),
        ("images", []),
        ("formulas", {"A1": "=1"}),
        ("page_setup", {"orientation": "portrait"}),
        ("visibility", "hidden"),
    ]:
        with pytest.raises(ValueError, match="grid"):
            _write([{"name": "S", "grid": np.array([[1.0]]), key: value}])


def test_table_name_width_and_overlap_validation():
    with pytest.raises(ValueError, match="table name"):
        _rows(tables=[{"name": "Bad Name", "ref": "A1:B2", "columns": ["A", "B"]}])
    with pytest.raises(ValueError, match="spans"):
        _rows(tables=[{"name": "T", "ref": "A1:C2", "columns": ["A", "B"]}])
    with pytest.raises(ValueError, match="overlaps"):
        _rows(
            tables=[
                {"name": "Table1", "ref": "A1:B2", "columns": ["A", "B"]},
                {"name": "Table2", "ref": "B2:C3", "columns": ["B", "C"]},
            ]
        )


@pytest.mark.parametrize(
    "options",
    [
        {"dv": [{"type": "whole", "sqref": "A1", "operator": "bogus"}]},
        {"dv": [{"type": "whole", "sqref": "A1", "errorStyle": "bogus"}]},
        {"dv": [{"type": "whole", "sqref": "A1", "imeMode": "bogus"}]},
        {"cell_styles": {(0, 0): {"alignment": {"vertical": "sideways"}}}},
        {
            "cell_styles": {
                (0, 0): {
                    "fill": {
                        "type": "gradient",
                        "degree": 999,
                        "stops": [
                            {"position": 0, "color": "FF0000"},
                            {"position": 1, "color": "00FF00"},
                        ],
                    }
                }
            }
        },
        {"page_setup": {"orientation": "sideways"}},
        {"page_setup": {"scale": 999}},
        {"page_setup": {"paper_size": 9999}},
        {"tab_color": "notacolor"},
    ],
)
def test_invalid_style_validation_and_page_fields_rejected(options):
    with pytest.raises(ValueError):
        _rows(**options)


def test_grid_limits_and_masked_arrays_rejected():
    np = pytest.importorskip("numpy")
    with pytest.raises(ValueError, match="limit"):
        _write([{"name": "S", "grid": np.zeros((2, 16385))}])
    with pytest.raises(ValueError, match="limit"):
        _write([{"name": "S", "grid": np.zeros((1048577, 1))}])
    masked = np.ma.array([[1.0, 2.0]], mask=[[False, True]])
    with pytest.raises(ValueError, match="masked"):
        _write([{"name": "S", "grid": masked}])


@pytest.mark.parametrize("row", ["abc", b"abc", {"a": 1}])
def test_bare_rows_rejected(row):
    with pytest.raises(ValueError, match="sequence"):
        _write([{"name": "S", "rows": [row]}])


def test_all_hidden_workbook_rejected():
    with pytest.raises(ValueError, match="visible"):
        _write(
            [
                {"name": "A", "rows": [[1]], "visibility": "hidden"},
                {"name": "B", "rows": [[2]], "visibility": "veryHidden"},
            ]
        )


def test_filter_names_and_active_tab_validation():
    with pytest.raises(ValueError, match="auto_filter"):
        _rows(auto_filter="not-a-range")
    with pytest.raises(ValueError, match="col_id"):
        _rows(auto_filter={"ref": "A1:B2", "columns": [{"col_id": 99}]})
    with pytest.raises(ValueError, match="defined name"):
        _write([{"name": "S", "rows": [[1]]}], defined_names=[{"name": "1Bad", "value": "S!$A$1"}])
    with pytest.raises(ValueError, match="defined name"):
        _write([{"name": "S", "rows": [[1]]}], defined_names=[{"name": "A1", "value": "S!$A$1"}])
    with pytest.raises(ValueError, match="local_sheet_id"):
        _write(
            [{"name": "S", "rows": [[1]]}],
            defined_names=[{"name": "N", "value": "S!$A$1", "local_sheet_id": 99}],
        )
    with pytest.raises(ValueError, match="active_tab"):
        _write([{"name": "S", "rows": [[1]]}], active_tab=99)


def test_drawing_comment_merge_and_hyperlink_validation():
    series = [{"val": "S!$A$1:$A$2"}]
    with pytest.raises(ValueError, match="color"):
        _rows(charts=[{"type": "line", "series": [{**series[0], "color": "not-a-color"}]}])
    with pytest.raises(ValueError, match="anchor"):
        _rows(charts=[{"type": "line", "anchor": "not-a-cell", "series": series}])
    with pytest.raises(ValueError, match="JPEG"):
        _rows(images=[{"data": bytes.fromhex("FFD8FFE000104A46494600010100000100010000FFDB00")}])
    with pytest.raises(ValueError, match="GIF"):
        _rows(images=[{"data": b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff\x2c"}])
    with pytest.raises(ValueError, match="duplicate comment"):
        _rows(comments=[{"ref": "A1", "text": "a"}, {"ref": "A1", "text": "b"}])
    with pytest.raises(ValueError, match="overlaps"):
        _rows(merges=["A1:B2", "B2:C3"])
    with pytest.raises(ValueError, match="unknown hyperlink key"):
        _rows(hyperlinks=[{"ref": "A1", "targt": "https://example.org"}])
    with pytest.raises(ValueError, match="target or location"):
        _rows(hyperlinks=[{"ref": "A1"}])


def test_control_characters_and_sheet_name_controls_rejected():
    with pytest.raises(ValueError, match="control character"):
        _rows(rows=[["a\x01b"]])
    with pytest.raises(Exception, match="sheet name"):
        _write([{"name": "a\tb", "rows": [[1]]}])


def test_editable_bounds_are_validated():
    wb = kyrax.Workbook()
    ws = wb.active
    ws["A1"] = 1
    assert list(ws.iter_rows(min_row=0, max_row=0, min_col=1, max_col=1, values_only=True)) == [
        (1,)
    ]
    with pytest.raises(ValueError):
        list(ws.iter_rows(min_row=-1))
    for call in (
        lambda: ws.insert_rows(1048577),
        lambda: ws.delete_rows(1048577),
        lambda: ws.insert_cols(16385),
        lambda: ws.delete_cols(16385),
    ):
        with pytest.raises(ValueError):
            call()


def test_pandas_bytesio_and_explicit_empty_strings(tmp_path):
    pd = pytest.importorskip("pandas")
    path = tmp_path / "pandas.xlsx"
    kyrax.write_excel_turbo(path, [{"name": "S", "rows": [["a", "b", "c"], [1, "", ""]]}])
    df = pd.read_excel(io.BytesIO(path.read_bytes()), engine="kyrax", na_filter=False)
    assert df.iloc[0, 1:].tolist() == ["", ""]


def test_json_and_csv_strictness(tmp_path):
    from kyrax import io as kio

    out = tmp_path / "out.xlsx"
    for payload in (b'[{"a":1}] trailing', b'[{"a":1},]', b'[{"a":1,"a":2}]'):
        with pytest.raises(Exception):
            kio.json_bytes_to_sheet(payload, out)
    with pytest.raises(Exception):
        kio.csv_bytes_to_sheet(b"a,b\n\xff\xfe,1\n", out)


def test_formula_dependencies_quote_sheet_names():
    from kyrax.formulas import dependencies

    assert dependencies("='My Sheet'!A1") == ["'My Sheet'!A1"]
    assert dependencies("='O''Brien'!A1") == ["'O''Brien'!A1"]


def test_stub_parses_and_has_no_duplicate_top_level_names():
    stub = Path(__file__).resolve().parents[2] / "python" / "kyrax" / "_kyrax.pyi"
    if not stub.exists():
        pytest.skip("repository stub not available from copied test tree")
    tree = ast.parse(stub.read_text(encoding="utf-8"))
    names = [node.name for node in tree.body if isinstance(node, (ast.ClassDef, ast.FunctionDef))]
    assert len(names) == len(set(names))


def test_existing_dv_cf_partial_delete_keeps_survivors(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    ws = original.active
    ws["A2"] = 1
    ws["A5"] = 5
    ws.add_data_validation(OpenpyxlDataValidation(type="whole", formula1="1", sqref="A2:A5"))
    ws.conditional_formatting.add("B2:B5", OpenpyxlFormulaRule(formula=["B2>0"]))
    original.save(source)

    wb = kyrax.load_workbook(source)
    wb.active.delete_rows(2)
    result = tmp_path / "result.xlsx"
    wb.save(result)

    sheet = openpyxl.load_workbook(result).active
    assert str(sheet.data_validations.dataValidation[0].sqref) == "A2:A4"
    assert str(list(sheet.conditional_formatting)[0].sqref) == "B2:B4"


def test_editable_cf_priority_is_preserved(tmp_path):
    wb = kyrax.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.conditional_formatting.add("A1:A2", Rule(type="expression", formula=["A1>0"], priority=7))
    out = tmp_path / "cf.xlsx"
    wb.save(out)
    rule = list(openpyxl.load_workbook(out).active.conditional_formatting)[0].rules[0]
    assert rule.priority == 7


def test_freeze_and_chartsheet_names_are_validated():
    with pytest.raises(ValueError, match="freeze_panes"):
        _rows(freeze_panes="not-a-cell")
    with pytest.raises(Exception, match="sheet name"):
        _write([{"name": "S", "rows": [[1]]}], chartsheets=[{"name": "Bad/Name"}])
    with pytest.raises(Exception, match="sheet name"):
        _write([{"name": "S", "rows": [[1]]}], chartsheets=[{"name": "S"}])
