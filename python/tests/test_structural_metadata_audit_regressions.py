import kyrax
import openpyxl
import pytest
from kyrax.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table


@pytest.mark.parametrize("operation", [
    lambda ws: ws.insert_rows(2),
    lambda ws: ws.move_range("C2:C2", rows=1),
])
def test_structural_edit_with_existing_comment_is_refused(tmp_path, operation):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active["C2"] = 3
    original.active["C2"].comment = Comment("note", "author")
    original.save(source)

    wb = kyrax.load_workbook(source)
    operation(wb.active)
    with pytest.raises(Exception, match="structural edits of existing comments"):
        wb.save(tmp_path / "out.xlsx")
    sheet = openpyxl.load_workbook(source).active
    assert sheet["C2"].value == 3
    assert sheet["C2"].comment.text == "note"


@pytest.mark.parametrize("operation", [
    lambda ws: ws.insert_cols(2),
    lambda ws: ws.delete_cols(2),
])
def test_existing_table_column_count_change_is_refused(tmp_path, operation):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.append(["first", "second", "third"])
    original.active.append([1, 2, 3])
    original.active.add_table(Table(displayName="Example", ref="A1:C2"))
    original.save(source)

    wb = kyrax.load_workbook(source)
    operation(wb.active)
    with pytest.raises(Exception, match="change the table's column count"):
        wb.save(tmp_path / "out.xlsx")
    table = openpyxl.load_workbook(source).active.tables["Example"]
    assert table.ref == "A1:C2"
    assert [col.name for col in table.tableColumns] == ["first", "second", "third"]


@pytest.mark.parametrize("metadata", ["merge", "hyperlink", "validation"])
def test_pending_metadata_moves_with_inserted_row(tmp_path, metadata):
    wb = kyrax.Workbook()
    ws = wb.active
    ws["A2"] = "value"
    if metadata == "hyperlink":
        ws["A2"].hyperlink = "https://example.com"
    elif metadata == "merge":
        ws.merge_cells("A2:B2")
    else:
        ws.add_data_validation(DataValidation(type="whole", sqref="A2", formula1="1"))
    ws.insert_rows(2)
    output = tmp_path / "out.xlsx"
    wb.save(output)

    sheet = openpyxl.load_workbook(output).active
    assert sheet["A3"].value == "value"
    if metadata == "hyperlink":
        assert sheet["A3"].hyperlink.target == "https://example.com"
        assert sheet["A2"].hyperlink is None
    elif metadata == "merge":
        assert {str(r) for r in sheet.merged_cells.ranges} == {"A3:B3"}
    else:
        assert [str(d.sqref) for d in sheet.data_validations.dataValidation] == ["A3"]
