import openpyxl
import pytest
import kyrax
from openpyxl.worksheet.table import Table


def test_unmerge_then_insert_keeps_merge_removed(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.merge_cells("A2:B2")
    original.save(source)
    wb = kyrax.load_workbook(source)
    ws = wb.active
    ws.unmerge_cells("A2:B2")
    ws.insert_rows(2)
    ws["B3"] = "visible"
    result = tmp_path / "result.xlsx"
    wb.save(result)
    sheet = openpyxl.load_workbook(result).active
    assert not sheet.merged_cells.ranges
    assert sheet["B3"].value == "visible"


def test_insert_then_unmerge_uses_shifted_existing_merge(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active["A2"] = "value"
    original.active.merge_cells("A2:B2")
    original.save(source)
    wb = kyrax.load_workbook(source)
    wb.active.insert_rows(2)
    wb.active.unmerge_cells("A3:B3")
    result = tmp_path / "result.xlsx"
    wb.save(result)
    assert not openpyxl.load_workbook(result).active.merged_cells.ranges


def test_move_intersecting_existing_table_refused(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.append(["one", "two"])
    original.active.append([1, 2])
    original.active.add_table(Table(displayName="Sample", ref="A1:B2"))
    original.save(source)
    wb = kyrax.load_workbook(source)
    wb.active.move_range("A1:B2", rows=2)
    with pytest.raises(Exception, match="intersect existing tables"):
        wb.save(tmp_path / "result.xlsx")


def test_move_away_from_existing_table_is_allowed(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.append(["one", "two"])
    original.active.append([1, 2])
    original.active.add_table(Table(displayName="Sample", ref="A1:B2"))
    original.active["D4"] = "outside"
    original.save(source)
    wb = kyrax.load_workbook(source)
    wb.active.move_range("D4:D4", rows=1)
    result = tmp_path / "result.xlsx"
    wb.save(result)
    assert openpyxl.load_workbook(result).active["D5"].value == "outside"
