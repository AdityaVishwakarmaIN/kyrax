import io

import kyrax
import openpyxl
import pytest


def write_sheet(**options):
    return kyrax.write_excel_turbo_bytes([{"name": "S", "rows": [[1], [2]], **options}])


@pytest.mark.parametrize("options", [
    {"dv": [{"type": "wholle", "sqref": "A1"}]},
    {"dv": [{"type": "whole", "sqref": "not-a-range"}]},
    {"cell_styles": {(0, 0): {"alignment": {"horizontal": "wrong"}}}},
    {"charts": [{"type": "line", "legend_pos": "wrong"}]},
    {"charts": [{"type": "line", "series": [{"val": "S!$A$1:$A$2", "marker": {"symbol": "circle", "size": 0}}]}]},
])
def test_reject_malformed_feature_inputs(options):
    with pytest.raises(ValueError):
        write_sheet(**options)


def test_valid_feature_inputs_round_trip():
    data = write_sheet(
        dv=[{"type": "whole", "sqref": "A1:A2"}],
        cell_styles={(0, 0): {"alignment": {"horizontal": "center"}}},
        charts=[{"type": "line", "legend_pos": "b", "series": [{"val": "S!$A$1:$A$2", "marker": {"symbol": "circle", "size": 7}}]}],
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.active["A1"].alignment.horizontal == "center"
    assert len(wb.active.data_validations.dataValidation) == 1
    assert len(wb.active._charts) == 1

@pytest.mark.parametrize("options", [
    {"merges": ["nonsense"]},
    {"hyperlinks": [{"ref": "nonsense", "target": "https://example.org"}]},
    {"comments": [{"ref": "nonsense", "text": "note"}]},
    {"tables": [{"name": "Sales", "ref": "garbage", "columns": ["A"]}]},
    {"tables": [{"name": "Sales", "ref": "A1:B2", "columns": ["A", "B"]}, {"name": "sales", "ref": "D1:E2", "columns": ["D", "E"]}]},
])
def test_reject_malformed_metadata(options):
    with pytest.raises(ValueError):
        write_sheet(**options)


def test_valid_metadata_round_trip():
    data = kyrax.write_excel_turbo_bytes([{"name": "S", "rows": [["A", "B"], [1, 2]],
        "tables": [{"name": "Sales", "ref": "A1:B2", "columns": ["A", "B"]}],
        "hyperlinks": [{"ref": "A2", "target": "https://example.org"}],
        "comments": [{"ref": "B2", "text": "note"}], "merges": ["D1:E1"]}])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert list(wb.active.tables) == ["Sales"]
    assert wb.active["A2"].hyperlink.target == "https://example.org"
    assert wb.active["B2"].comment.text == "note"
    assert "D1:E1" in wb.active.merged_cells


def test_reject_duplicate_table_names_across_sheets():
    table = {"name": "Sales", "ref": "A1:A2", "columns": ["A"]}
    with pytest.raises(ValueError, match="duplicate table"):
        kyrax.write_excel_turbo_bytes([{"name": "S1", "rows": [["A"], [1]], "tables": [table]},
                                      {"name": "S2", "rows": [["A"], [2]], "tables": [table]}])


def test_reject_truncated_png_and_accept_complete_image():
    import base64

    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")
    for end in range(8, len(png)):
        with pytest.raises(ValueError):
            write_sheet(images=[{"data": png[:end]}])
    workbook = openpyxl.load_workbook(io.BytesIO(write_sheet(images=[{"data": png}])))
    assert len(workbook.active._images) == 1
