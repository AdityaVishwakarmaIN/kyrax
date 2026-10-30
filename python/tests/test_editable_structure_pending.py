import openpyxl
import pytest
import kyrax


@pytest.mark.parametrize('operation,expected', [
    (lambda s: s.insert_rows(2, 2), {'A1': 'one', 'A4': 'two', 'B5': 'three'}),
    (lambda s: s.delete_rows(2), {'A1': 'one', 'B2': 'three'}),
    (lambda s: s.insert_cols(2, 2), {'A1': 'one', 'A2': 'two', 'D3': 'three'}),
    (lambda s: s.delete_cols(1), {'A3': 'three'}),
    (lambda s: s.move_range('A2:B3', rows=3, cols=2), {'A1': 'one', 'C5': 'two', 'D6': 'three'}),
])
def test_new_sheet_structural_edits(tmp_path, operation, expected):
    wb = kyrax.Workbook()
    ws = wb.active
    ws.append(['one'])
    ws.append(['two'])
    ws.append([None, 'three'])
    operation(ws)
    path = tmp_path / 'new.xlsx'
    wb.save(str(path))
    actual = openpyxl.load_workbook(path).active
    assert {cell.coordinate: cell.value for row in actual for cell in row if cell.value is not None} == expected


def test_pending_edits_follow_operations_in_order(tmp_path):
    path = tmp_path / 'existing.xlsx'
    wb = kyrax.Workbook()
    wb.active['A2'] = 'old'
    wb.active['B3'] = 'overwritten'
    wb.save(str(path))
    wb = kyrax.load_workbook(str(path))
    ws = wb.active
    ws['A2'] = '=B2+$B$1'
    ws['B3'] = 'pending destination'
    ws.move_range('A2:A2', rows=1, cols=1, translate=True)
    ws.insert_rows(2)
    ws['A2'] = 'later edit'
    wb.save(str(path))
    actual = openpyxl.load_workbook(path).active
    assert actual['A2'].value == 'later edit'
    assert actual['A3'].value is None
    assert actual['B4'].value == '=C3+$B$1'
