"""Independent readback of editable styles and attached validations."""

import kyrax
from kyrax.styles import Alignment, Protection
from kyrax.worksheet.datavalidation import DataValidation
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook


def test_alignment_and_protection_survive_save(tmp_path):
    # Exercise both the fresh workbook and the append-only style merge.
    source = tmp_path / "source.xlsx"
    OpenpyxlWorkbook().save(source)
    for i, workbook in enumerate((kyrax.Workbook(), kyrax.load_workbook(source))):
        sheet = workbook.active
        sheet["A1"] = "both"
        sheet["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
        sheet["A1"].protection = Protection(locked=False, hidden=True)
        output = tmp_path / f"styles{i}.xlsx"
        workbook.save(output)
        cell = load_workbook(output).active["A1"]
        assert (cell.alignment.horizontal, cell.alignment.wrap_text) == ("center", True)
        assert (cell.protection.locked, cell.protection.hidden) == (False, True)


def test_attached_validation_changes_and_options_survive_save(tmp_path):
    workbook = kyrax.Workbook()
    dv = DataValidation(type="whole", formula1="1", sqref="A1", errorStyle="stop", imeMode="on")
    workbook.active.add_data_validation(dv)
    dv.formula1 = "5"
    dv.allow_blank = True
    dv.errorStyle = "warning"
    dv.imeMode = "off"
    output = tmp_path / "validation.xlsx"
    workbook.save(output)
    actual = load_workbook(output).active.data_validations.dataValidation[0]
    assert (actual.formula1, actual.allowBlank, actual.errorStyle, actual.imeMode) == (
        "5",
        True,
        "warning",
        "off",
    )
    # Loading and modifying validations must retain the two attributes.
    reopened = kyrax.load_workbook(output)
    reopened.active.add_data_validation(DataValidation(type="whole", sqref="B1"))
    reopened.save(output)
    actual = load_workbook(output).active.data_validations.dataValidation[0]
    assert (actual.errorStyle, actual.imeMode) == ("warning", "off")
