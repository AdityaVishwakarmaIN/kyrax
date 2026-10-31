import kyrax.pandas  # noqa: F401 - register the engine
import openpyxl
import pandas as pd
import pytest

pytestmark = pytest.mark.skipif(
    int(pd.__version__.split(".")[0]) < 3,
    reason="DataFrame.to_excel(autofilter=...) requires pandas 3.0+",
)


def test_pandas_autofilter_single_sheet(tmp_path):
    path = tmp_path / "single.xlsx"
    df = pd.DataFrame({"name": ["a", "b", "c"], "value": [1, 2, 3]})
    df.to_excel(path, engine="kyrax", autofilter=True)

    ws = openpyxl.load_workbook(path).active
    assert ws.auto_filter.ref == "A1:C4"


def test_pandas_autofilter_multiple_sheets(tmp_path):
    path = tmp_path / "multiple.xlsx"
    first = pd.DataFrame({"name": ["a", "b"], "value": [1, 2]})
    second = pd.DataFrame({"city": ["x", "y", "z"]})

    with pd.ExcelWriter(path, engine="kyrax") as writer:
        first.to_excel(writer, sheet_name="First", autofilter=True)
        second.to_excel(writer, sheet_name="Second", autofilter=True)

    workbook = openpyxl.load_workbook(path)
    assert workbook["First"].auto_filter.ref == "A1:C3"
    assert workbook["Second"].auto_filter.ref == "A1:B4"


def test_pandas_autofilter_startrow_startcol(tmp_path):
    path = tmp_path / "offset.xlsx"
    df = pd.DataFrame({"name": ["a", "b", "c"], "value": [1, 2, 3]})
    df.to_excel(path, engine="kyrax", startrow=2, startcol=1, autofilter=True)

    ws = openpyxl.load_workbook(path).active
    assert ws.auto_filter.ref == "B3:D6"


def test_pandas_autofilter_disabled_leaves_no_filter(tmp_path):
    path = tmp_path / "disabled.xlsx"
    pd.DataFrame({"name": ["a"], "value": [1]}).to_excel(path, engine="kyrax")

    ws = openpyxl.load_workbook(path).active
    assert ws.auto_filter.ref is None
